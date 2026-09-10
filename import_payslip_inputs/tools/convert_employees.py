# -*- coding: utf-8 -*-
"""Converter v3: อ่านไฟล์ HR ฉบับเดียว (Rev.04+) ที่มี 2 ชีต -> ชุดไฟล์ import (data/ready3)
    Sheet "P'หมูอัพเดท"        = ผังองค์กร (authoritative: ใครอยู่ในระบบ) + วันเริ่มงาน + สถานะว่าจ้าง
    Sheet "Data Odoo(ต้องเพิ่ม)" = ข้อมูลส่วนตัว (จับคู่ด้วยรหัส)
รัน:  PYTHONUTF8=1 python custom_addons/import_payslip_inputs/tools/convert_employees.py
กติกาเดิม: ข้อมูลผิด -> ว่าง + FIX_REPORT | แผนก 4 ชั้น | คำนำหน้าแยกแต่ต้น | marital ว่าง = single"""
import csv
import datetime
import hashlib
import os
import re

import openpyxl
from openpyxl.styles import Font, PatternFill

SRC = r"d:\odoo18\other\สรุปรายชื่อส่งให้พี่เอ็ก IT Up to Odoo Rev.04 เพิ่มรักษาการ.xlsx"
SHEET_ORG = "P'หมูอัพเดท"
SHEET_PERSONAL = 'Data Odoo(ต้องเพิ่ม)'
STATE_CSV = r"d:\odoo18\odoo\addons\base\data\res.country.state.csv"
OUT = r"d:\odoo18\custom_addons\import_payslip_inputs\data\ready3"
os.makedirs(OUT, exist_ok=True)

TH_PROVINCE = {
    'กรุงเทพ': 'Bangkok', 'กรุงเทพฯ': 'Bangkok', 'กรุงเทพมหานคร': 'Bangkok',
    'อำนาจเจริญ': 'Amnat Charoen', 'อ่างทอง': 'Ang Thong', 'บึงกาฬ': 'Bueng Kan',
    'บุรีรัมย์': 'Buriram', 'ฉะเชิงเทรา': 'Chachoengsao', 'ชัยนาท': 'Chai Nat',
    'ชัยภูมิ': 'Chaiyaphum', 'จันทบุรี': 'Chanthaburi', 'เชียงใหม่': 'Chiang Mai',
    'เชียงราย': 'Chiang Rai', 'ชลบุรี': 'Chonburi', 'ชุมพร': 'Chumphon',
    'กาฬสินธุ์': 'Kalasin', 'กำแพงเพชร': 'Kamphaeng Phet', 'กาญจนบุรี': 'Kanchanaburi',
    'ขอนแก่น': 'Khon Kaen', 'กระบี่': 'Krabi', 'ลำปาง': 'Lampang', 'ลำพูน': 'Lamphun',
    'เลย': 'Loei', 'ลพบุรี': 'Lopburi', 'แม่ฮ่องสอน': 'Mae Hong Son',
    'มหาสารคาม': 'Maha Sarakham', 'มุกดาหาร': 'Mukdahan', 'นครนายก': 'Nakhon Nayok',
    'นครปฐม': 'Nakhon Pathom', 'นครพนม': 'Nakhon Phanom',
    'นครราชสีมา': 'Nakhon Ratchasima', 'โคราช': 'Nakhon Ratchasima',
    'นครสวรรค์': 'Nakhon Sawan', 'นครศรีธรรมราช': 'Nakhon Si Thammarat',
    'น่าน': 'Nan', 'นราธิวาส': 'Narathiwat', 'หนองบัวลำภู': 'Nong Bua Lamphu',
    'หนองคาย': 'Nong Khai', 'นนทบุรี': 'Nonthaburi', 'ปทุมธานี': 'Pathum Thani',
    'ปัตตานี': 'Pattani', 'พังงา': 'Phang Nga', 'พัทลุง': 'Phatthalung',
    'พะเยา': 'Phayao', 'เพชรบูรณ์': 'Phetchabun', 'เพชรบุรี': 'Phetchaburi',
    'พิจิตร': 'Phichit', 'พิษณุโลก': 'Phitsanulok',
    'พระนครศรีอยุธยา': 'Phra Nakhon Si Ayutthaya', 'อยุธยา': 'Phra Nakhon Si Ayutthaya',
    'แพร่': 'Phrae', 'ภูเก็ต': 'Phuket', 'ปราจีนบุรี': 'Prachinburi',
    'ประจวบคีรีขันธ์': 'Prachuap Khiri Khan', 'ระนอง': 'Ranong', 'ราชบุรี': 'Ratchaburi',
    'ระยอง': 'Rayong', 'ร้อยเอ็ด': 'Roi Et', 'สระแก้ว': 'Sa Kaeo',
    'สกลนคร': 'Sakon Nakhon', 'สมุทรปราการ': 'Samut Prakan', 'สมุทรสาคร': 'Samut Sakhon',
    'สมุทรสงคราม': 'Samut Songkhram', 'สระบุรี': 'Saraburi', 'สตูล': 'Satun',
    'สิงห์บุรี': 'Sing Buri', 'ศรีสะเกษ': 'Sisaket', 'สงขลา': 'Songkhla',
    'สุโขทัย': 'Sukhothai', 'สุพรรณบุรี': 'Suphan Buri', 'สุราษฎร์ธานี': 'Surat Thani',
    'สุรินทร์': 'Surin', 'ตาก': 'Tak', 'ตรัง': 'Trang', 'ตราด': 'Trat',
    'อุบลราชธานี': 'Ubon Ratchathani', 'อุดรธานี': 'Udon Thani',
    'อุทัยธานี': 'Uthai Thani', 'อุตรดิตถ์': 'Uttaradit', 'ยะลา': 'Yala',
    'ยโสธร': 'Yasothon',
    # คำสะกดผิดที่เคยพบ
    'ศรีษะเกษ': 'Sisaket', 'ศรีษาะเกษ': 'Sisaket', 'ปราจีน': 'Prachinburi',
    'สิุรินทร์': 'Surin', 'สุราษฏร์ธานี': 'Surat Thani', 'กาฬสินธ์': 'Kalasin',
}
COUNTRY = {
    'ไทย': 'base.th', 'ประเทศไทย': 'base.th', 'thailand': 'base.th',
    'เมียนมา': 'base.mm', 'เมียนมาร์': 'base.mm', 'พม่า': 'base.mm', 'myanmar': 'base.mm',
    'เขมร': 'base.kh', 'กัมพูชา': 'base.kh', 'cambodia': 'base.kh',
    'ลาว': 'base.la', 'laos': 'base.la',
}
GENDER = {'ชาย': 'male', 'หญิง': 'female'}
MARITAL = {'โสด': 'single', 'สมรส': 'married', 'ไม่จดทะเบียน': 'cohabitant'}
CERT = {
    'ประถมศึกษาปีที่ 3': 'primary_3', 'ประถมศึกษาปีที่ 6': 'primary_6',
    'มัธยมศึกษาปีที่ 3': 'secondary_3', 'มัธยมศึกษาปีที่ 6': 'secondary_6',
    'กศน': 'nfe', 'ปวช': 'vocational_cert', 'ปวท': 'technical_cert',
    'ปวส': 'high_vocational', 'ปริญญาตรี': 'bachelor', 'ปริญญาโท': 'master',
    'ปริญญาเอก': 'doctor',
    # การสะกดแบบอื่นที่พบใน Rev.04
    'มัธยมศึกษาตอนต้น': 'secondary_3', 'มัธยมศึกษาตอนปลาย': 'secondary_6',
    'มัธยมศึกษาต้นปลาย': 'secondary_6', 'มศ.5': 'secondary_6', 'มศ.เอ็น5': 'secondary_6',
    'ประถมศึกษา (ป.6)': 'primary_6', 'ประถมศึกษา(ป.6)': 'primary_6',
    'ประถมศึกษาตอนปลาย': 'primary_6', 'กศน.ตอนปลาย': 'nfe', 'กศน.ตอนต้น': 'nfe',
}
CORP_EMAIL = {
    '10000': 'thawach.sa@autozonegroup.in.th', '10001': 'pavana.sa@autozonegroup.in.th',
    '10005': 'ubol.ea@autozonegroup.in.th', '10025': 'suntaree.ka@autozonegroup.in.th',
    '10045': 'prapaporn.le@autozonegroup.in.th', '10069': 'kittipong.ch@autozonegroup.in.th',
    '10403': 'wanna.pi@autozonegroup.in.th', '10464': 'jintana.ka@autozonegroup.in.th',
    '11968': 'naroj.sa@autozonegroup.in.th', '12697': 'jeenapa.pa@autozonegroup.in.th',
    '13596': 'nuttana.na@autozonegroup.in.th', '13800': 'siriyakorn.ka@autozonegroup.in.th',
    '13936': 'thanisak.mo@autozonegroup.in.th', '14109': 'pornpawee.yi@autozonegroup.in.th',
    '14190': 'naruemon.ph@autozonegroup.in.th',
}
USER_LINKS = [
    ('10000', 'Thawach Sawasdiworn'), ('10001', 'Pavana Sawasdiworn'),
    ('10005', 'Ubol Eamwichien'), ('10025', 'Suntaree Kaewmeesee'),
    ('10045', 'Prapaporn Lertngamdee'), ('10069', 'Kittipong Chernrungroj'),
    ('10403', 'Wanna Pimsim'), ('10464', 'Jintana Katasiy'),
    ('11968', 'Naroj Salee'), ('12697', 'Jeenapa Panyathip'),
    ('13596', 'Nuttana Naruephatpalangkorn'), ('13800', 'Siriyakorn Kanbut'),
    ('13936', 'Thanisak Moonsiri'), ('14109', 'Pornpawee Yindeesuk'),
    ('14190', 'Naruemon Phadungthaitham'), ('14382', 'จิตตรา วิษาจันทร์'),
    ('14387', 'วรรณ์นิพา มณีมงคลแก้ว'), ('12157', 'Suphawat Pornbonkul'),
]
TITLE_RE = re.compile(r'^(นางสาว|น\.ส\.|นาย|นาง|คุณ|MRS\.?|MISS\s?|MR\.?|MS\.?)\s*', re.IGNORECASE)
T_NORM = {'น.ส.': 'นางสาว', 'MR': 'Mr.', 'MR.': 'Mr.', 'MRS': 'Mrs.', 'MRS.': 'Mrs.',
          'MS': 'Ms.', 'MS.': 'Ms.', 'MISS': 'Miss'}
COMBINING = '\u0e31\u0e34\u0e35\u0e36\u0e37\u0e38\u0e39\u0e3a\u0e47\u0e48\u0e49\u0e4a\u0e4b\u0e4c\u0e4d\u0e4e'

STATE_XMLID = {}
with open(STATE_CSV, encoding='utf-8') as f:
    for row in csv.DictReader(f):
        if row['country_id:id'] == 'base.th':
            STATE_XMLID[row['name']] = 'base.' + row['id']


def dedup_marks(s):
    out = []
    for ch in s:
        if out and ch == out[-1] and ch in COMBINING:
            continue
        out.append(ch)
    return ''.join(out)


def clean(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = str(int(v))
    s = str(v).strip()
    return '' if s in ('*', '-') else s


def spaces(s):
    return re.sub(r'\s+', ' ', s).strip()


def split_name(full):
    s = spaces(dedup_marks(full)).rstrip('-').strip()
    m = TITLE_RE.match(s)
    if not m:
        return '', s
    raw = m.group(1).strip()
    return T_NORM.get(raw.upper(), T_NORM.get(raw, raw)), s[m.end():].strip()


def norm_person(s):
    return split_name(s)[1].casefold()


def fix_phone(s):
    d = re.sub(r'\D', '', s)
    if not d:
        return ''
    if len(d) == 11 and d.startswith('66'):
        d = '0' + d[2:]
    if len(d) in (8, 9) and not d.startswith('0'):
        d = '0' + d
    return d


def to_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        y = v.year - 543 if v.year > 2100 else v.year
        try:
            return datetime.date(y, v.month, v.day), None
        except ValueError as e:
            return None, str(e)
    s = clean(v)
    if not s:
        return None, None
    m = re.match(r'^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$', s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m2 = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
        if not m2:
            return None, 'รูปแบบวันที่ไม่รู้จัก'
        y, mo, d = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
    if y > 2100:
        y -= 543
    try:
        return datetime.date(y, mo, d), None
    except ValueError as e:
        return None, str(e)


def dept_xmlid(path):
    return 'dept_' + hashlib.md5('|'.join(path).encode()).hexdigest()[:10]


# ---------------- อ่านไฟล์ ----------------
wb = openpyxl.load_workbook(SRC, data_only=True)

ws1 = wb[SHEET_ORG]
# คอลัมน์ Sheet1: 2 รหัส, 3 ชื่อ, 4 ตำแหน่ง, 5 JobLevel, 8 หัวหน้า, 10 สังกัด(สาขา),
#                11 กลุ่มสายงาน, 12 งาน, 13 หน่วย, 14 แผนก, 15 ฝ่าย, 16 สถานะว่าจ้าง, 17 วันเริ่มงาน
org = {}
seats = []   # แถวซ้ำของรหัสเดิม = "เก้าอี้รักษาการ" (Rev.04: HeadCount 1 / OnHand 0)
report = []
for r in range(2, ws1.max_row + 1):
    def c1(col):
        return clean(ws1.cell(row=r, column=col).value)
    code = c1(2)
    if not code:
        continue
    row = {'code': code, 'name': c1(3), 'job': c1(4), 'level': c1(5), 'manager': c1(8),
           'branch': c1(10), 'group': c1(11), 'work': c1(12), 'unit': c1(13),
           'dept': c1(14), 'division': c1(15), 'hire_status': c1(16),
           'start_raw': ws1.cell(row=r, column=17).value}
    if code in org:
        seats.append(row)
        continue
    org[code] = row

ws2 = wb[SHEET_PERSONAL]
# คอลัมน์ Sheet2 (ข้าม 2 แถวหัว): 1 รหัส, 10 email, 11 phone, 15 street, 16 city, 17 state,
#   18 zip, 19 pphone, 20 idcard, 21 gender, 22 birthday, 23 pob, 24 cob, 25 bank,
#   26 econtact, 27 ephone, 28 marital, 29 cert, 30 sfield, 31 sschool, 32-35 visa/permit
personal = {}
for r in range(3, ws2.max_row + 1):
    code = clean(ws2.cell(row=r, column=1).value)
    if code:
        personal[code] = {i: ws2.cell(row=r, column=i).value for i in range(1, 36)}

print(f"ผัง (Sheet1): {len(org)} คน | ข้อมูลส่วนตัว (Sheet2): {len(personal)} คน")

# ---------------- ผังแผนก 4 ชั้น (รวม path ของแถวเก้าอี้รักษาการด้วย) ----------------
dept_nodes = {}
for o in list(org.values()) + seats:
    path = []
    for name in (o['division'], o['dept'], o['unit'], o['work']):
        if not name:
            continue
        path.append(name)
        key = tuple(path)
        if key not in dept_nodes:
            dept_nodes[key] = (dept_xmlid(path), name,
                               dept_xmlid(path[:-1]) if len(path) > 1 else '')

# ---------------- ประกอบร่างพนักงาน ----------------
branches = set()
employees = []
for code, o in org.items():
    p = personal.get(code)
    title, name = split_name(o['name'])
    e = {'id': f'emp_{code}', 'registration_number': code, 'name': name,
         'title_th': title, 'job_id': spaces(o['job']), 'job_level': o['level'],
         'category_ids': ','.join(x for x in (o['group'], o['hire_status']) if x)}

    def rep(field, val, action='ปรับเป็นค่าว่าง'):
        report.append((code, name, field, str(val), action))

    levels = [x for x in (o['division'], o['dept'], o['unit'], o['work']) if x]
    e['department_id/id'] = dept_xmlid(levels) if levels else ''
    branch = o['branch'].upper()
    e['work_location_id/id'] = f'wl_{branch}' if branch else ''
    if branch:
        branches.add(branch)

    # วันเริ่มงาน (เก็บไว้ทำสัญญาจ้าง ไม่ import ลง employee)
    sd, err = to_date(o['start_raw'])
    if err or (sd and sd.year < 1980):
        rep('วันเริ่มงาน', o['start_raw'], f'ต้องตรวจ ({err or "ปีเก่าผิดปกติ"})')
    e['_start'] = sd.isoformat() if sd else ''

    if p:
        def pc(i):
            return clean(p[i])
        e['work_email'] = CORP_EMAIL.get(code) or pc(10)
        e['work_phone'] = fix_phone(pc(11))
        e['private_street'] = pc(15)
        e['private_city'] = pc(16)
        state = dedup_marks(pc(17))
        en = TH_PROVINCE.get(state) or (state if state in STATE_XMLID else None)
        if state and not (en and en in STATE_XMLID):
            rep('จังหวัด', state)
        e['private_state_id/id'] = STATE_XMLID.get(en, '') if en else ''
        e['private_zip'] = pc(18)
        e['private_country_id/id'] = 'base.th' if (e['private_street'] or e['private_state_id/id']) else ''
        e['private_phone'] = fix_phone(pc(19))
        idcard = re.sub(r'\D', '', pc(20))
        if idcard and len(idcard) != 13:
            rep('เลขบัตรประชาชน', pc(20))
            idcard = ''
        e['identification_id'] = idcard
        g = dedup_marks(pc(21))
        e['gender'] = GENDER.get(g, '')
        if g and not e['gender']:
            rep('เพศ', g)
        bd, err = to_date(p[22])
        if err:
            rep('วันเกิด', p[22], f'ปรับเป็นค่าว่าง ({err})')
            bd = None
        if bd and not (1946 <= bd.year <= 2011):
            rep('วันเกิด', bd.isoformat(), 'ปรับเป็นค่าว่าง (อายุไม่สมเหตุผล)')
            bd = None
        e['birthday'] = bd.isoformat() if bd else ''
        e['place_of_birth'] = pc(23)
        cob = dedup_marks(pc(24)).lower()
        e['country_of_birth/id'] = COUNTRY.get(cob, '')
        if cob and not e['country_of_birth/id']:
            rep('ประเทศเกิด', pc(24))
        e['emergency_contact'] = spaces(pc(26))
        e['emergency_phone'] = fix_phone(pc(27))
        mar = pc(28)
        e['marital'] = MARITAL.get(mar, '') or 'single'
        if mar and mar not in MARITAL:
            rep('สถานะสมรส', mar, 'ปรับเป็น single')
        cert = dedup_marks(pc(29)).rstrip('.')
        e['certificate'] = CERT.get(cert, '')
        if cert and not e['certificate']:
            rep('วุฒิการศึกษา', pc(29))
        e['study_field'] = pc(30)
        e['study_school'] = pc(31)
        for key, i in (('visa_no', 32), ('permit_no', 33)):
            v = p[i]
            if isinstance(v, (datetime.datetime, datetime.date)):
                rep(key, v, 'ปรับเป็นค่าว่าง (กรอกเป็นวันที่)')
                e[key] = ''
            else:
                e[key] = clean(v)
        for key, i in (('visa_expire', 34), ('work_permit_expiration_date', 35)):
            d, err = to_date(p[i])
            if err:
                rep(key, p[i], f'ปรับเป็นค่าว่าง ({err})')
            e[key] = d.isoformat() if d else ''
        e['_bank'] = re.sub(r'\D', '', pc(25))
        e['_nickname'] = pc(3)
    else:
        for k in ('work_email', 'work_phone', 'private_street', 'private_city',
                  'private_state_id/id', 'private_zip', 'private_country_id/id',
                  'private_phone', 'identification_id', 'gender', 'birthday',
                  'place_of_birth', 'country_of_birth/id', 'emergency_contact',
                  'emergency_phone', 'certificate', 'study_field', 'study_school',
                  'visa_no', 'permit_no', 'visa_expire', 'work_permit_expiration_date'):
            e[k] = ''
        e['marital'] = 'single'
        e['_bank'] = ''
        e['_nickname'] = ''
        rep('ข้อมูลส่วนตัว', '(ไม่อยู่ในชีต Data Odoo)', 'รอ HR ส่งเพิ่ม')
    employees.append(e)

# ---------------- เก้าอี้รักษาการ (จากแถวซ้ำใน Sheet1) ----------------
seat_by_owner = {}
managers = []
for i, s in enumerate(seats, 1):
    title, base_name = split_name(s['name'])
    mlabel = re.search(r'\(([^)]*)\)', s['job'])
    label = mlabel.group(1) if mlabel else s['job']
    sid = f"emp_seat_{s['code']}_{i}"
    levels = [x for x in (s['division'], s['dept'], s['unit'], s['work']) if x]
    branch = s['branch'].upper()
    if branch:
        branches.add(branch)
    rec = {'id': sid, 'registration_number': '', 'name': f'{base_name} ({label})',
           'title_th': title, 'job_id': spaces(s['job']), 'job_level': s['level'],
           'category_ids': ','.join(x for x in (s['group'], 'ตำแหน่งรักษาการ') if x),
           'department_id/id': dept_xmlid(levels) if levels else '',
           'work_location_id/id': f'wl_{branch}' if branch else '',
           'marital': 'single', '_start': '', '_bank': '', '_nickname': ''}
    for k in ('work_email', 'work_phone', 'private_street', 'private_city',
              'private_state_id/id', 'private_zip', 'private_country_id/id',
              'private_phone', 'identification_id', 'gender', 'birthday',
              'place_of_birth', 'country_of_birth/id', 'emergency_contact',
              'emergency_phone', 'certificate', 'study_field', 'study_school',
              'visa_no', 'permit_no', 'visa_expire', 'work_permit_expiration_date'):
        rec[k] = ''
    employees.append(rec)
    managers.append((sid, f"emp_{s['code']}"))   # เก้าอี้ขึ้นกับตัวจริง
    seat_by_owner.setdefault(s['code'], []).append({'sid': sid, 'division': s['division']})
    report.append((s['code'], rec['name'], 'เก้าอี้รักษาการ', s['job'],
                   f'สร้างกล่องเก้าอี้ {sid} (ไม่มีสัญญาจ้าง/ไม่เข้า payroll)'))

# ---------------- หัวหน้างาน (route เข้าเก้าอี้ถ้าสังกัดฝ่ายเดียวกับเก้าอี้) ----------------
by_norm = {}
for e in employees:
    by_norm.setdefault(e['name'].casefold(), []).append(e['id'])
for code, o in org.items():
    m = o['manager']
    if not m:
        report.append((code, org[code]['name'], 'หัวหน้างาน', '(ว่าง)', 'ข้าม'))
        continue
    if norm_person(m) == norm_person(o['name']):
        continue
    hits = by_norm.get(norm_person(m), [])
    if len(hits) != 1:
        report.append((code, org[code]['name'], 'หัวหน้างาน', m,
                       'ข้าม (ซ้ำหลายคน)' if hits else 'ข้าม (ไม่พบในไฟล์)'))
        continue
    boss_id = hits[0]
    boss_code = boss_id.replace('emp_', '')
    cands = [st for st in seat_by_owner.get(boss_code, [])
             if st['division'] and st['division'] == o['division']
             and o['division'] != org.get(boss_code, {}).get('division')]
    if len(cands) == 1:
        boss_id = cands[0]['sid']   # ลูกทีมฝ่ายเดียวกับเก้าอี้ -> ขึ้นกับเก้าอี้
    elif len(cands) > 1:
        report.append((code, org[code]['name'], 'หัวหน้างาน', m,
                       'เก้าอี้รักษาการหลายตัวในฝ่ายเดียวกัน - ผูกกับตัวจริงไว้ก่อน'))
    managers.append((f'emp_{code}', boss_id))

# ---------------- เขียนไฟล์ ----------------
H_FILL = PatternFill('solid', fgColor='1F4E78')
H_FONT = Font(color='FFFFFF', bold=True)
RED = PatternFill('solid', fgColor='C00000')
TEXT_COLS = {'registration_number', 'identification_id', 'work_phone',
             'private_phone', 'emergency_phone', 'private_zip', 'รหัสพนักงาน'}


def write_xlsx(fname, headers, data_rows, red_cols=()):
    wb2 = openpyxl.Workbook()
    ws = wb2.active
    ws.title = 'Data'
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = H_FILL
        cell.font = H_FONT
    for col in red_cols:
        ws.cell(row=1, column=headers.index(col) + 1).fill = RED
    text_idx = [i for i, h in enumerate(headers) if h in TEXT_COLS]
    for row in data_rows:
        ws.append(row)
        for i in text_idx:
            ws.cell(row=ws.max_row, column=i + 1).number_format = '@'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(13, min(40, len(str(h)) + 6))
    ws.freeze_panes = 'A2'
    wb2.save(os.path.join(OUT, fname))
    print(f"  {fname}: {len(data_rows)} แถว")


dept_rows = [[x[0], x[1], x[2]] for x in (dept_nodes[k] for k in sorted(dept_nodes, key=len))]
write_xlsx('00a_departments.xlsx', ['id', 'name', 'parent_id/id'], dept_rows)
write_xlsx('00b_work_locations.xlsx', ['id', 'name', 'address_id/id'],
           [[f'wl_{b}', b, 'base.main_partner'] for b in sorted(branches)])

EMP_HEADERS = ['id', 'registration_number', 'name', 'title_th', 'job_id', 'job_level',
               'category_ids', 'department_id/id', 'work_location_id/id',
               'work_email', 'work_phone', 'private_street', 'private_city',
               'private_state_id/id', 'private_zip', 'private_country_id/id',
               'private_phone', 'identification_id', 'gender', 'birthday',
               'place_of_birth', 'country_of_birth/id', 'emergency_contact',
               'emergency_phone', 'marital', 'certificate', 'study_field',
               'study_school', 'visa_no', 'permit_no', 'visa_expire',
               'work_permit_expiration_date']
write_xlsx('01_employees_READY.xlsx', EMP_HEADERS,
           [[e[h] for h in EMP_HEADERS] for e in employees])
write_xlsx('02_managers.xlsx', ['id', 'parent_id/id'], [list(m) for m in managers])

links = [[f'emp_{c}', u] for c, u in USER_LINKS if c in org]
for c, u in USER_LINKS:
    if c not in org:
        report.append((c, u, 'ผูก user', '(ไม่อยู่ในผังใหม่)', 'ข้าม'))
write_xlsx('03_link_users.xlsx', ['id', 'user_id'], links)

# 04: เหลือกรอกแค่เงินเดือน (วันเริ่มงานมาจากไฟล์ HR แล้ว)
emp_by_code = {e['registration_number']: e for e in employees}
rows04 = []
for code in sorted(org, key=lambda c: (org[c]['branch'], c)):
    e = emp_by_code[code]
    o = org[code]
    rows04.append([code, o['name'], o['job'], o['branch'], o['group'], o['hire_status'],
                   e['_start'], None])
write_xlsx('04_contract_data_to_fill.xlsx',
           ['รหัสพนักงาน', 'ชื่อ-นามสกุล', 'ตำแหน่ง', 'สาขา', 'กลุ่มสายงาน',
            'สถานะว่าจ้าง', 'วันเริ่มงาน (จากไฟล์ HR)', 'เงินเดือน (กรอก)'],
           rows04, red_cols=['เงินเดือน (กรอก)'])

write_xlsx('90_bank_accounts_later.xlsx', ['registration_number', 'name', 'เลขบัญชี'],
           [[e['registration_number'], e['name'], e['_bank']] for e in employees if e['_bank']])
write_xlsx('FIX_REPORT.xlsx', ['registration_number', 'ชื่อ', 'ข้อมูล', 'ค่าเดิม', 'การจัดการ'],
           [list(x) for x in report])

from collections import Counter
print(f"\nสรุป: พนักงาน {len(employees)} | แผนก {len(dept_nodes)} | สาขา {len(branches)}"
      f" | หัวหน้า {len(managers)} | user {len(links)} | FIX {len(report)}")
print("กลุ่มสายงาน:", dict(Counter(o['group'] for o in org.values())))
print("สถานะว่าจ้าง:", dict(Counter(o['hire_status'] for o in org.values())))
print("วันเริ่มงานใช้ได้:", sum(1 for e in employees if e['_start']), "/", len(employees))
print("\nFIX แยกประเภท:")
for k, n in Counter(x[2] for x in report).most_common():
    print(f"  {n:3}  {k}")
