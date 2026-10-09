# -*- coding: utf-8 -*-
"""นำเข้าจากไฟล์ Excel ที่ Export ออกจากใบนับแล้วกรอกกลับมา (อ่านตามชื่อหัวคอลัมน์ ไม่ยึดตำแหน่ง)

* ชีต "ใบตรวจนับ": หัวตาราง = แถวที่มี "Material"/"รหัสสินค้า" — อ่าน ยอดตรวจนับได้ (หรือ ยืนยันจำนวน) / หมายเหตุ / ID   แถวที่ยืนยันจำนวนว่าง = ยังไม่นับ ข้าม
* ชีต "กระทบยอด" (ถ้ามี): หัวตาราง = แถวที่มี "+ บวกรับ" + "- หักจ่าย" + "สถานะ" — ค่าที่ต่างจากในใบจะเขียนทับ
  qty_in/qty_out (ติ๊ก คีย์เอง ให้)
* ชีต "หมายเหตุผลต่าง" (ถ้ามี): หัวตาราง = แถวที่มี "สาเหตุของผลต่าง" — อ่าน สาเหตุ / เอกสาร /
  ผู้รับผิดชอบ / การดำเนินการ / วันที่แล้วเสร็จ   เขียนเฉพาะช่องที่ไม่ว่าง
จับคู่สินค้า: ID (product id) → รหัสสินค้า → ชื่อสินค้า
รองรับไฟล์รุ่นแรก (คอลัมน์ ลำดับ/รหัสสินค้า/ชื่อสินค้า/หน่วย/ยอดคงเหลือ/ยอดตรวจนับ/ผลต่าง/หมายเหตุ/ID) ด้วย
"""
import base64
import io
from datetime import date, datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

# ชื่อหัวคอลัมน์ที่ยอมรับ (เทียบแบบ startswith หลังตัดช่องว่าง)
COLS_COUNT = {
    "id": ("ID",),
    "code": ("Material", "รหัสสินค้า"),
    "name": ("Description", "ชื่อสินค้า"),
    "counted": ("ยืนยันจำนวน", "ยอดตรวจนับ"),
    "note": ("หมายเหตุ",),
}
COLS_RECON = {
    "code": ("Material", "รหัสสินค้า"),
    "name": ("Description", "ชื่อสินค้า"),
    "qty_in": ("+ บวกรับ", "บวกรับ"),
    "qty_out": ("- หักจ่าย", "หักจ่าย"),
    "status": ("สถานะ",),
}
COLS_NOTES = {
    "code": ("Material", "รหัสสินค้า"),
    "name": ("Description", "ชื่อสินค้า"),
    "reason": ("สาเหตุของผลต่าง", "สาเหตุ"),
    "ref_doc": ("เอกสาร",),
    "responsible": ("ผู้รับผิดชอบ",),
    "action_taken": ("การดำเนินการ",),
    "date_resolved": ("วันที่แล้วเสร็จ",),
}


class StockCountImport(models.TransientModel):
    _name = "az.stock.count.import"
    _description = "นำเข้ายอดนับ / สาเหตุผลต่าง จาก Excel"

    count_id = fields.Many2one("az.stock.count", required=True, ondelete="cascade")
    file = fields.Binary("ไฟล์ Excel (.xlsx)", required=True)
    filename = fields.Char()
    overwrite = fields.Boolean(
        "ทับยอดที่คีย์ไว้แล้ว", default=True,
        help="ไม่ติ๊ก = รายการที่คีย์ยอดนับในจอไปแล้วจะไม่ถูกแก้ด้วยค่าจากไฟล์ (สาเหตุผลต่างทับเสมอถ้าไฟล์มีค่า)",
    )
    result = fields.Text("ผลการนำเข้า", readonly=True)

    @api.model
    def _to_float(self, value):
        if value is None or value == "":
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            raise UserError(_("ค่า '%s' ไม่ใช่ตัวเลข", value))

    @api.model
    def _to_date(self, value):
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @api.model
    def _find_header(self, ws, spec, must):
        """หาแถวหัวตาราง → (row_idx, {key: col_idx}) หรือ (None, None)"""
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            cols = {}
            for c_idx, val in enumerate(row):
                if not isinstance(val, str):
                    continue
                text = val.strip()
                for key, names in spec.items():
                    if key not in cols and any(text.startswith(n) for n in names):
                        cols[key] = c_idx
                        break
            if must in cols and ("code" in cols or "id" in cols):
                return r_idx, cols
        return None, None

    def _match_line(self, maps, row, cols):
        by_id, by_code, by_name = maps
        line = None
        if "id" in cols and row[cols["id"]] not in (None, ""):
            try:
                line = by_id.get(int(float(row[cols["id"]])))
            except (TypeError, ValueError):
                line = None
        if not line and "code" in cols and row[cols["code"]] not in (None, ""):
            line = by_code.get(str(row[cols["code"]]).strip())
        if not line and "name" in cols and row[cols["name"]]:
            line = by_name.get(str(row[cols["name"]]).strip())
        return line

    def action_import(self):
        self.ensure_one()
        count = self.count_id
        count._check_state(("counting", "counted"), "นำเข้าจาก Excel")
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_("เซิร์ฟเวอร์ไม่มี openpyxl"))
        try:
            wb = load_workbook(io.BytesIO(base64.b64decode(self.file)), data_only=True, read_only=True)
        except Exception as e:
            raise UserError(_("เปิดไฟล์ไม่ได้ (ต้องเป็น .xlsx): %s", e))

        lines = count.line_ids
        by_id = {l.product_id.id: l for l in lines}
        by_code, by_name = {}, {}
        for l in lines:
            if l.product_id.default_code:
                by_code.setdefault(l.product_id.default_code.strip(), l)
            by_name.setdefault((l.product_id.name or "").strip(), l)
        maps = (by_id, by_code, by_name)

        summary = []
        problems = []
        count_done = notes_done = recon_done = False
        for ws in wb.worksheets:
            # ---- ชีตกระทบยอด: รับเข้า/จ่ายออก ที่แก้เอง ----
            if not recon_done:
                h, cols = self._find_header(ws, COLS_RECON, "qty_in")
                if h is not None and "qty_out" in cols and "status" in cols:
                    recon_done = True
                    updated = unmatched = 0
                    for row in ws.iter_rows(min_row=h + 2, values_only=True):
                        row = list(row) + [None] * (max(cols.values()) + 1 - len(row))
                        if row[cols["code"]] in (None, ""):
                            continue
                        line = self._match_line(maps, row, cols)
                        if not line:
                            unmatched += 1
                            continue
                        vals = {}
                        for key in ("qty_in", "qty_out"):
                            v = self._to_float(row[cols[key]])
                            if v is not None and abs(v - getattr(line, key)) > 0.00001:
                                vals[key] = v
                        if vals:
                            line.write(vals)
                            updated += 1
                    summary.append(_("ชีต '%(s)s': แก้รับเข้า/จ่ายออกตามไฟล์ %(u)d รายการ, จับคู่สินค้าไม่ได้ %(x)d",
                                     s=ws.title, u=updated, x=unmatched))
                    continue
            # ---- ชีตยอดนับ ----
            if not count_done:
                h, cols = self._find_header(ws, COLS_COUNT, "counted")
                if h is not None:
                    count_done = True
                    updated = skipped = unmatched = 0
                    for row in ws.iter_rows(min_row=h + 2, values_only=True):
                        row = list(row) + [None] * (max(cols.values()) + 1 - len(row))
                        if not any(row[c] not in (None, "") for k, c in cols.items() if k in ("id", "code")):
                            continue  # แถวว่าง / หัวหมวด / ลายเซ็น
                        line = self._match_line(maps, row, cols)
                        if not line:
                            if row[cols["counted"]] not in (None, ""):
                                unmatched += 1
                                if len(problems) < 20:
                                    problems.append("%s | %s" % (row[cols.get("code", 0)] or "", row[cols.get("name", 0)] or ""))
                            continue
                        vals = {}
                        if "note" in cols and row[cols["note"]] not in (None, ""):
                            vals["note"] = str(row[cols["note"]]).strip()
                        qty = self._to_float(row[cols["counted"]])
                        if qty is None:
                            skipped += 1
                        elif line.counted and not self.overwrite and count.state == "counting":
                            skipped += 1
                        elif count.state == "counting" or self.overwrite:
                            vals.update({"qty_counted": qty, "counted": True})
                            updated += 1
                        if vals:
                            line.write(vals)
                    summary.append(_("ชีต '%(s)s': นำเข้ายอดนับ %(u)d รายการ, ข้าม (ไม่ได้กรอก/ไม่ทับ) %(k)d, จับคู่สินค้าไม่ได้ %(x)d",
                                     s=ws.title, u=updated, k=skipped, x=unmatched))
                    continue
            # ---- ชีตสาเหตุผลต่าง ----
            if not notes_done:
                h, cols = self._find_header(ws, COLS_NOTES, "reason")
                if h is not None:
                    notes_done = True
                    updated = unmatched = 0
                    for row in ws.iter_rows(min_row=h + 2, values_only=True):
                        row = list(row) + [None] * (max(cols.values()) + 1 - len(row))
                        vals = {}
                        for key in ("reason", "ref_doc", "responsible", "action_taken"):
                            if key in cols and row[cols[key]] not in (None, ""):
                                vals[key] = str(row[cols[key]]).strip()
                        if "date_resolved" in cols:
                            d = self._to_date(row[cols["date_resolved"]])
                            if d:
                                vals["date_resolved"] = d
                        if not vals or row[cols["code"]] in (None, ""):
                            continue  # แถวว่าง / ลายเซ็น
                        line = self._match_line(maps, row, cols)
                        if not line:
                            unmatched += 1
                            if len(problems) < 20:
                                problems.append("%s | %s" % (row[cols.get("code", 0)] or "", row[cols.get("name", 0)] or ""))
                            continue
                        line.write(vals)
                        updated += 1
                    summary.append(_("ชีต '%(s)s': นำเข้าสาเหตุผลต่าง %(u)d รายการ, จับคู่สินค้าไม่ได้ %(x)d",
                                     s=ws.title, u=updated, x=unmatched))
        if not count_done and not notes_done and not recon_done:
            raise UserError(_("ไม่พบหัวตารางที่รู้จักในไฟล์ (ต้องมีคอลัมน์ Material/รหัสสินค้า และ ยอดตรวจนับได้ / บวกรับ-หักจ่าย / สาเหตุของผลต่าง)"))
        text = "\n".join(summary)
        if problems:
            text += "\n" + _("จับคู่ไม่ได้:") + "\n" + "\n".join(problems)
        count.message_post(body=_("นำเข้าจากไฟล์ %(f)s:<br/>%(s)s", f=self.filename or "", s=text.replace("\n", "<br/>")))
        self.result = text
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }
